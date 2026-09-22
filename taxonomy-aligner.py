"""
Universal QA Toolkit — Intra-Franchise Taxonomy Aligner (taxonomy-aligner.py)

1. Consistency: within each franchise, rows whose subcategory differs from the
   brand's majority are classified before anything is changed:
     realign            clear majority, same storefront name -> safe to align
     allowlisted        brand listed in --allowlist (intentionally mixed formats,
                        e.g. a DIY chain with garden centres) -> never changed
     format_variant     the outlier rows use a different business_name than the
                        majority rows ('Petrol Car Wash' vs 'Petrol') -> review
     mixed_review       tie, weak majority (< --min-share) or small group
                        (< --min-rows) -> review
   Only 'realign' rows are changed, and only when --output is given.
2. Reference taxonomy (--taxonomy): every category/subcategory is checked for an
   exact match against the official list (unknown subcategory, subcategory under
   the wrong category, case/spacing-only mismatch).

--output keeps every sheet, formula and format of the input workbook: only the
realigned cells change. Nothing is written without --output.
"""

import argparse
import os
import sys
from collections import defaultdict

import pandas as pd

import qa_common as qc

REALIGN, ALLOWLISTED, FORMAT_VARIANT, MIXED = "realign", "allowlisted", "format_variant", "mixed_review"


# --------------------------------------------------------------------------
# Reference taxonomy
# --------------------------------------------------------------------------

def load_taxonomy(path):
    """Return {subcategory: {category, ...}} from a CSV/XLSX with Category and
    Subcategory columns (header row found automatically, merged category cells
    forward-filled, all sheets read) or a JSON list of {"name"|"subcategory",
    "category"} objects (the converter's seed format)."""
    if not os.path.exists(path):
        print(f"Error: Taxonomy file '{path}' not found.", file=sys.stderr)
        sys.exit(1)
    pairs = []
    if path.lower().endswith(".json"):
        for item in qc.load_json(path):
            sub = (item.get("subcategory") or item.get("name") or "").strip()
            cat = (item.get("category") or item.get("category_name") or "").strip()
            if sub and cat:
                pairs.append((cat, sub))
    else:
        if path.lower().endswith(".csv"):
            sheets = {"csv": pd.read_csv(path, header=None, dtype=str)}
        else:
            sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=str)
        for raw in sheets.values():
            pairs.extend(_pairs_from_grid(raw.fillna("")))
    if not pairs:
        print(f"Error: No category/subcategory pairs found in '{path}'.", file=sys.stderr)
        sys.exit(1)
    taxonomy = defaultdict(set)
    for cat, sub in pairs:
        taxonomy[sub].add(cat)
    qc.LOG.info("Reference taxonomy: %d subcategories in %d categories",
                len(taxonomy), len({c for c, _ in pairs}))
    return taxonomy


def _pairs_from_grid(grid):
    for r in range(min(len(grid), 15)):
        cells = [str(v).strip().lower() for v in grid.iloc[r]]
        if "category" in cells and "subcategory" in cells:
            cat_i, sub_i = cells.index("category"), cells.index("subcategory")
            break
    else:
        return []
    pairs, current = [], ""
    for _, row in grid.iloc[r + 1:].iterrows():
        cat, sub = str(row.iloc[cat_i]).strip(), str(row.iloc[sub_i]).strip()
        if cat.startswith("▼"):  # section banner rows
            continue
        current = cat or current
        if sub and current:
            pairs.append((current, sub))
    return pairs


def check_against_taxonomy(df, taxonomy, category_col, subcat_col):
    by_norm = defaultdict(list)
    for sub in taxonomy:
        by_norm[qc.norm_text(sub)].append(sub)
    findings = []
    for (cat, sub), idx in df.groupby([category_col, subcat_col], sort=False).groups.items():
        cat, sub = cat.strip(), sub.strip()
        if not sub:
            issue, detail = "subcategory_missing", ""
        elif sub in taxonomy:
            if cat in taxonomy[sub]:
                continue
            issue, detail = "subcategory_under_wrong_category", f"belongs to {sorted(taxonomy[sub])}"
        elif qc.norm_text(sub) in by_norm:
            issue, detail = "subcategory_spelling_differs", f"official spelling {by_norm[qc.norm_text(sub)]}"
        else:
            issue, detail = "subcategory_not_in_taxonomy", "the import will not find this subcategory"
        for i in idx:
            findings.append({"sheet": df.at[i, qc.SHEET_COL], "row": int(df.at[i, qc.ROW_COL]),
                             "category": cat, "subcategory": sub, "issue": issue, "detail": detail})
    return findings


# --------------------------------------------------------------------------
# Intra-franchise consistency
# --------------------------------------------------------------------------

def classify_outliers(df, franchise_col, subcat_col, name_col, allowlist, min_share, min_rows):
    allowed = {qc.norm_text(b) for b in allowlist}
    has_names = name_col in df.columns
    keys = df[franchise_col].map(qc.norm_text)
    results = []
    for key, group in df[keys != ""].groupby(keys[keys != ""], sort=False):
        subcats = group[subcat_col].str.strip()
        subcats = subcats[subcats != ""]
        counts = subcats.value_counts()
        if len(counts) < 2:
            continue
        brand = group[franchise_col].str.strip().value_counts().index[0]
        majority, top = counts.index[0], int(counts.iloc[0])
        share = top / len(subcats)
        tie = int(counts.iloc[1]) == top
        majority_names = set(group.loc[subcats[subcats == majority].index, name_col].map(qc.norm_text)) if has_names else set()
        for i in subcats[subcats != majority].index:
            if key in allowed:
                cls, why = ALLOWLISTED, "brand is on the mixed-format allowlist"
            elif tie or share < min_share or len(subcats) < min_rows:
                cls, why = MIXED, ("tie between subcategories" if tie else
                                   f"weak majority ({share:.0%})" if share < min_share else
                                   f"only {len(subcats)} rows")
            elif has_names and qc.norm_text(group.at[i, name_col]) not in majority_names:
                cls, why = FORMAT_VARIANT, f"storefront name '{group.at[i, name_col]}' differs from the majority rows"
            else:
                cls, why = REALIGN, "clear majority, same storefront name"
            results.append({
                "sheet": group.at[i, qc.SHEET_COL],
                "row_number": int(group.at[i, qc.ROW_COL]),
                "franchise_name": brand,
                "current_subcategory": subcats[i],
                "expected_majority_subcategory": majority,
                "majority_share": f"{top}/{len(subcats)} branches",
                "classification": cls,
                "reason": why,
            })
    return results


def _write_realigned(excel_path, output_path, subcat_col, changes):
    """Copy the workbook and change only the realigned cells, keeping every sheet."""
    if excel_path.lower().endswith(".csv"):
        df = pd.read_csv(excel_path, dtype=str, keep_default_na=False)
        for c in changes:
            df.at[c["row_number"] - 2, subcat_col] = c["expected_majority_subcategory"]
        df.to_csv(output_path, index=False)
        return
    import openpyxl
    wb = openpyxl.load_workbook(excel_path)
    columns = {}
    for c in changes:
        ws = wb[c["sheet"]]
        if ws.title not in columns:
            header = [str(v.value).strip() if v.value is not None else "" for v in ws[1]]
            columns[ws.title] = header.index(subcat_col) + 1
        ws.cell(row=c["row_number"], column=columns[ws.title], value=c["expected_majority_subcategory"])
    wb.save(output_path)


def align_franchise_taxonomy(excel_path, franchise_col="franchise_name", subcat_col="subcategory", output_path=None,
                             allowlist=(), min_share=0.6, min_rows=3, name_col="business_name", report=None,
                             taxonomy_path=None, category_col="category", sheet=None, engine=None):
    if output_path and os.path.abspath(output_path) == os.path.abspath(excel_path):
        print("Error: --output must be a different file from the input.", file=sys.stderr)
        sys.exit(1)
    with qc.timed("Load"):
        df = qc.load_dataset(excel_path, sheet=sheet, required_columns=[franchise_col, subcat_col], engine=engine)

    print(f"Analyzing intra-franchise subcategory consistency for "
          f"{df[franchise_col].map(qc.norm_text).replace('', pd.NA).nunique():,} unique franchise brands...")
    with qc.timed("Consistency"):
        results = classify_outliers(df, franchise_col, subcat_col, name_col, allowlist, min_share, min_rows)
    by_class = defaultdict(list)
    for r in results:
        by_class[r["classification"]].append(r)

    print("\n" + "=" * 60)
    print("TAXONOMY ALIGNMENT AUDIT RESULTS")
    print("=" * 60)
    print(f"Franchise Brands Audited       : {df[franchise_col].map(qc.norm_text).replace('', pd.NA).nunique():,}")
    print(f"Subcategory Outlier Rows Found : {len(results):,}")
    print(f"  realign (safe, clear majority) : {len(by_class[REALIGN]):,}")
    print(f"  format_variant (review)        : {len(by_class[FORMAT_VARIANT]):,}")
    print(f"  mixed_review (review)          : {len(by_class[MIXED]):,}")
    print(f"  allowlisted (no action)        : {len(by_class[ALLOWLISTED]):,}")
    print("=" * 60)

    for cls in (REALIGN, FORMAT_VARIANT, MIXED, ALLOWLISTED):
        if by_class[cls]:
            print(f"\nSample [{cls}]:")
            for r in by_class[cls][:10]:
                print(f"  {r['sheet']}!{r['row_number']} [{r['franchise_name']}]: '{r['current_subcategory']}' "
                      f"--> '{r['expected_majority_subcategory']}' ({r['majority_share']}; {r['reason']})")

    taxonomy_findings = []
    if taxonomy_path:
        if category_col not in df.columns:
            qc.LOG.warning("Column '%s' not found: taxonomy pairs cannot be checked", category_col)
        else:
            with qc.timed("Reference taxonomy"):
                taxonomy_findings = check_against_taxonomy(df, load_taxonomy(taxonomy_path), category_col, subcat_col)
            counts = defaultdict(int)
            for f in taxonomy_findings:
                counts[f["issue"]] += 1
            print(f"\nReference taxonomy check: {len(taxonomy_findings):,} row(s) {dict(counts) or '(all match)'}")
            seen = set()
            for f in taxonomy_findings:
                key = (f["issue"], f["category"], f["subcategory"])
                if key not in seen and len(seen) < 15:
                    seen.add(key)
                    print(f"  [{f['issue']}] {f['category']} / {f['subcategory']} {f['detail']}")

    if report:
        qc.write_records(results + [{**f, "classification": "taxonomy_" + f["issue"]} for f in taxonomy_findings], report)
    if output_path:
        _write_realigned(excel_path, output_path, subcat_col, by_class[REALIGN])
        print(f"\nSaved {len(by_class[REALIGN]):,} realigned cell(s) to: {output_path} "
              f"(other sheets and formatting kept; review classes left unchanged)")
    return results, taxonomy_findings


def main():
    parser = argparse.ArgumentParser(description="Intra-Franchise Subcategory Majority Taxonomy Aligner")
    parser.add_argument("excel_path", help="Path to input Excel dataset")
    parser.add_argument("--franchise-col", default="franchise_name", help="Franchise name column (default: franchise_name)")
    parser.add_argument("--subcat-col", default="subcategory", help="Subcategory column (default: subcategory)")
    parser.add_argument("--category-col", default="category", help="Category column (default: category)")
    parser.add_argument("--name-col", default="business_name", help="Storefront name column (default: business_name)")
    parser.add_argument("--output", help="Optional output Excel file path to save cleaned dataset")
    parser.add_argument("--allowlist", help="File of brands with intentionally mixed formats (one per line)")
    parser.add_argument("--min-share", type=float, default=0.6, help="Majority share needed to realign (default: 0.6)")
    parser.add_argument("--min-rows", type=int, default=3, help="Brand rows needed to realign (default: 3)")
    parser.add_argument("--taxonomy", help="Reference taxonomy (.xlsx/.csv with Category and Subcategory, or seed .json)")
    parser.add_argument("--report", help="Write every outlier and taxonomy finding to this .csv or .json file")
    parser.add_argument("--sheet", default=None, help="Only this sheet (default: every sheet with the columns)")
    parser.add_argument("--engine", default=None, help="Excel reader engine (default: calamine if installed)")
    qc.add_logging_args(parser)
    args = parser.parse_args()
    qc.setup_logging(args)

    align_franchise_taxonomy(args.excel_path, args.franchise_col, args.subcat_col, args.output,
                             allowlist=qc.load_list_file(args.allowlist), min_share=args.min_share,
                             min_rows=args.min_rows, name_col=args.name_col, report=args.report,
                             taxonomy_path=args.taxonomy, category_col=args.category_col,
                             sheet=args.sheet, engine=args.engine)


if __name__ == "__main__":
    main()
